"""舊車出廠年月改成直接打 6 碼數字 YYYYMM（使用者 2026-10-22 T-93）。

資料庫維持日期欄位，存成當月 1 日、顯示回 YYYYMM；手機跳數字鍵盤，格式錯誤即時提示且不送出。
匯出改成同樣的 6 碼，匯入讀得回來；匯入修正畫面也接受 6 碼。
"""
import datetime
from io import BytesIO
from unittest import mock

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook

from sales.forms import YearMonthDateField, YearMonthDateInput, year_month_hint
from sales.tests import test_old_vehicle_section as old_vehicle

TODAY = datetime.date(2026, 10, 22)


@mock.patch("sales.forms.timezone.localdate", return_value=TODAY)
class YearMonthDateFieldTests(SimpleTestCase):
    def test_six_digits_store_first_of_month(self, _today):
        field = YearMonthDateField(required=False)
        self.assertEqual(field.clean("201309"), datetime.date(2013, 9, 1))
        self.assertEqual(field.clean("202610"), datetime.date(2026, 10, 1))
        self.assertIsNone(field.clean(""))
        # 有分隔符號或貼上完整日期也收，只取年月。
        for text in ("2013/09", "2013-9", "2013年9月", "2013-09-01", "2013/9/15"):
            self.assertEqual(field.clean(text), datetime.date(2013, 9, 1), text)

    def test_invalid_values_explain_what_to_fix(self, _today):
        field = YearMonthDateField(required=False)
        cases = {"20139": "6 碼", "abcdef": "6 碼", "201313": "月份", "201300": "月份",
                 "194912": "年份", "202711": "年份", "202611": "晚於本月"}
        for text, expected in cases.items():
            with self.assertRaises(ValidationError, msg=text) as raised:
                field.clean(text)
            self.assertIn(expected, raised.exception.messages[0], text)

    def test_widget_shows_yyyymm_with_numeric_keyboard_and_hint(self, _today):
        html = YearMonthDateInput().render("operations-old_vehicle_manufactured_on", datetime.date(2013, 9, 1),
                                           attrs={"id": "id_ym"})
        self.assertIn('value="201309"', html)
        self.assertIn('inputmode="numeric"', html)
        self.assertIn('maxlength="6"', html)
        self.assertIn('data-year-month', html)
        self.assertIn('data-no-recent-values', html)
        self.assertIn('aria-describedby="id_ym_ym"', html)
        self.assertIn('id="id_ym_ym"', html)
        self.assertIn("2013 年 9 月・民國 102 年", html)
        self.assertIn("6 碼數字，例 201309", YearMonthDateInput().render("x", None, attrs={"id": "id_x"}))

    def test_hint_matches_javascript_wording(self, _today):
        self.assertEqual(year_month_hint("201309"), "2013 年 9 月・民國 102 年")
        self.assertEqual(year_month_hint(""), "6 碼數字，例 201309")


class OldVehicleYearMonthPageTests(old_vehicle.OldVehicleSectionTests.__bases__[0]):
    setUp = old_vehicle.OldVehicleSectionTests.setUp
    make_order = old_vehicle.OldVehicleSectionTests.make_order
    login_admin = old_vehicle.OldVehicleSectionTests.login_admin

    def save(self, order, value):
        profile = order.operations
        profile.refresh_from_db()
        return self.client.post(reverse("order_operations", args=[order.pk]), {
            "_section": "old_vehicle", "operations-financial_revision": profile.updated_at.isoformat(),
            "operations-old_vehicle_manufactured_on": value,
        }, HTTP_X_ORDER_WORKSPACE="1")

    def test_subsidy_step_shows_yyyymm_input(self):
        self.login_admin()
        order, _vehicle = self.make_order()
        order.operations.old_vehicle_manufactured_on = datetime.date(2013, 9, 1)
        order.operations.save(update_fields=["old_vehicle_manufactured_on"])
        html = self.client.get(reverse("order_detail", args=[order.pk])).content.decode()
        start = html.index('name="operations-old_vehicle_manufactured_on"')
        tag = html[html.rindex("<input", 0, start):html.index(">", start) + 1]
        self.assertIn('value="201309"', tag)
        self.assertIn('inputmode="numeric"', tag)
        self.assertNotIn('type="date"', tag)
        self.assertIn('form="ops-old-vehicle-form"', tag)
        self.assertIn(">舊車出廠年月<", html)
        self.assertIn("js/year-month-input.js", html)

    def test_autosave_stores_first_of_month_and_rejects_bad_month(self):
        self.login_admin()
        order, _vehicle = self.make_order()
        self.save(order, "201309")
        order.operations.refresh_from_db()
        self.assertEqual(order.operations.old_vehicle_manufactured_on, datetime.date(2013, 9, 1))
        response = self.save(order, "201313")
        self.assertEqual(response.json()["errors"]["operations-old_vehicle_manufactured_on"], ["月份要在 01 到 12 之間。"])
        order.operations.refresh_from_db()
        self.assertEqual(order.operations.old_vehicle_manufactured_on, datetime.date(2013, 9, 1))
        self.save(order, "")
        order.operations.refresh_from_db()
        self.assertIsNone(order.operations.old_vehicle_manufactured_on)

    def test_standalone_operations_page_uses_same_input(self):
        self.login_admin()
        order, _vehicle = self.make_order()
        order.operations.old_vehicle_manufactured_on = datetime.date(2005, 9, 1)
        order.operations.save(update_fields=["old_vehicle_manufactured_on"])
        html = self.client.get(reverse("order_operations", args=[order.pk])).content.decode()
        self.assertIn('value="200509"', html)
        self.assertIn("2005 年 9 月・民國 94 年", html)

    def test_export_writes_yyyymm(self):
        self.login_admin()
        order, _vehicle = self.make_order()
        order.operations.old_vehicle_manufactured_on = datetime.date(2013, 9, 1)
        order.operations.save(update_fields=["old_vehicle_manufactured_on"])
        sheet = load_workbook(BytesIO(self.client.get(reverse("operations_report_export")).content)).active
        headers = [cell.value for cell in sheet[1]]
        column = headers.index("出廠日期")
        values = [row[column] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertIn("201309", values)

