from datetime import date
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

from django.contrib.auth import get_user_model
from django.http import QueryDict
from django.test import TestCase, RequestFactory
from django.urls import reverse
from openpyxl import load_workbook

from sales.models import SalesOrder, SalesSource, VehicleModel, VehicleColor, PaymentRecord, OrderOperationsProfile
from sales.services.dashboard_metrics import build_dashboard_metrics, _percent_change
from sales.services.order_sorting import sort_context, sort_orders
from sales.services.sales_metrics import summarize_sales, filter_payment_risk
from sales.views import _operations_analysis, _operations_report_queryset


class CommandCenterTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_superuser('command-admin', password='test-only')
        cls.model = VehicleModel.objects.create(brand='測試', name='戰情測試機種', energy_type='gas')
        cls.color = VehicleColor.objects.create(vehicle_model=cls.model, name='白')

    def make_order(self, day=None, status='completed', price=80000, profit=20000, established=None):
        order = SalesOrder.objects.create(owner_name='戰情合成測試', owner_phone='0912345678',
            owner_address='合成地址', owner_id_number='A123456789', vehicle_model=self.model, color=self.color,
            registration_date=day, established_on=established or day, status=status, vehicle_price=price)
        OrderOperationsProfile.objects.filter(order=order).update(actual_disbursement=price,
            vehicle_cost=price-profit, legacy_finance_reconciliation={})
        return SalesOrder.objects.select_related('operations', 'vehicle_model', 'source').get(pk=order.pk)

    def test_same_period_excludes_future_and_cancelled(self):
        current = self.make_order(date(2026, 9, 10))
        self.make_order(date(2026, 8, 10), price=40000)
        self.make_order(date(2026, 8, 20))
        self.make_order(date(2026, 9, 20))
        for status in ('cancelled', 'cancel_refund_pending'):
            self.make_order(date(2026, 9, 10), status=status)
        metrics = build_dashboard_metrics(date(2026, 9, 12))
        self.assertEqual(metrics['performance']['count'], 1)
        self.assertEqual(metrics['performance']['sales_total'], current.vehicle_price)
        self.assertEqual(metrics['performance']['sales_change'], 100)
        self.assertEqual(metrics['period']['previous_end'], date(2026, 8, 12))
        self.assertEqual(metrics['new_orders'], 1)
        self.assertEqual(len(metrics['trend']), 12)
        self.assertEqual(metrics['trend'][-1]['count'], 1)
        self.assertEqual(metrics['trend'][-2]['count'], 2)
        self.assertEqual(len(metrics['charts']), 3)

    def test_month_boundary_short_month_and_zero_base(self):
        for today, previous in [(date(2026,3,30),date(2026,2,28)), (date(2024,3,30),date(2024,2,29)),
                                (date(2026,2,28),date(2026,1,31)), (date(2026,1,1),date(2025,12,1))]:
            with self.subTest(today=today):
                self.assertEqual(build_dashboard_metrics(today)['period']['previous_end'], previous)
        self.assertIsNone(_percent_change(100, 0))
        self.assertIsNone(_percent_change(0, 0))
        self.assertEqual(_percent_change(-50, -100), 50)

    def test_profit_completeness_and_shared_summary(self):
        ready = self.make_order(date(2026,9,1), profit=-100)
        pending = self.make_order(date(2026,9,2))
        OrderOperationsProfile.objects.filter(order=pending).update(legacy_finance_reconciliation={'status':'mismatch'})
        cancelled = self.make_order(date(2026,9,3), status='cancelled', profit=9000)
        rows = SalesOrder.objects.select_related('operations', 'vehicle_model').prefetch_related('payment_records')
        summary, models = _operations_analysis(rows)
        dashboard = build_dashboard_metrics(date(2026,9,12))
        self.assertEqual(summary['count'], 2)
        self.assertEqual(summary['net_profit'], -100)
        self.assertEqual(summary['profit_pending'], 1)
        self.assertEqual(models[0]['net_profit'], -100)
        self.assertEqual(dashboard['performance']['profit_total'], summary['net_profit'])
        self.assertEqual(dashboard['performance']['sales_total'], summary['vehicle_sales'])
        self.assertEqual(cancelled.operations.net_profit, 9000)
        self.assertEqual(ready.operations.net_profit, -100)
        empty = summarize_sales([cancelled])
        self.assertIsNone(empty['average_profit'])
        self.assertEqual(build_dashboard_metrics(date(2025,1,1))['charts'][2]['points'], [])

    def test_risk_counts_match_drilldown_and_do_not_duplicate_orders(self):
        outstanding = self.make_order(date(2026,9,1))
        refund = self.make_order(status='cancel_refund_pending')
        cancelled = self.make_order(status='cancelled')
        for order in (outstanding, cancelled):
            for _ in range(2):
                PaymentRecord.objects.create(order=order, expected_amount=500, received_amount=100, confirmed=False)
        metrics = build_dashboard_metrics(date(2026,9,12))
        self.assertEqual(metrics['risk'], {'outstanding':1,'unconfirmed':1,'refund':1})
        for risk, expected in [('outstanding',outstanding), ('unconfirmed',outstanding), ('refund',refund)]:
            query = _operations_report_queryset(RequestFactory().get('/', {'risk':risk}))
            self.assertEqual(list(query.values_list('pk', flat=True)), [expected.pk])
            self.assertEqual(query.count(), metrics['risk'][risk])

    def test_cancelled_display_export_and_original_money_preserved(self):
        self.client.force_login(self.user)
        for status in ('cancelled', 'cancel_refund_pending'):
            order = self.make_order(status=status, profit=12345)
            response = self.client.get(reverse('order_list'), {'status':status})
            self.assertNotContains(response, '待補領牌日')
            self.assertNotContains(response, '未領牌')
            self.assertNotContains(response, '12,345')
            self.assertContains(response, '取消訂單不列入成交淨利')
            self.assertContains(response, f'status-pill status-{status}')
            self.assertEqual(SalesOrder.objects.get(pk=order.pk).operations.net_profit, 12345)
        response = self.client.get(reverse('operations_report'), {'include_cancelled':1})
        self.assertEqual(response.context['analysis_summary']['count'], 0)
        self.assertEqual(response.context['page_obj'].paginator.count, 2)
        exported = self.client.get(reverse('operations_report_export'), {'include_cancelled':1})
        sheet = load_workbook(BytesIO(exported.content)).active
        column = [c.value for c in sheet[1]].index('單筆淨利')+1
        self.assertIsNone(sheet.cell(2,column).value)
        self.assertIsNone(sheet.cell(3,column).value)

    def test_source_fallback_and_sort_use_yusheng_without_renaming_people(self):
        order = self.make_order()
        self.assertEqual(order.source_display, '馭盛')
        self.assertEqual(sort_orders(SalesOrder.objects.all(), ['source']).get()._sort_source, '馭盛')
        staff = SalesSource.objects.create(name='其他門市承辦人', source_type='store')
        order.source = staff
        self.assertEqual(order.source_display, staff.name)
        order.registration_date = date(2026,9,1)
        order.status = 'cancelled'
        self.assertEqual(order.registration_display, '2026/09/01')

    def test_cancelled_profit_sort_is_empty_not_hidden(self):
        active = self.make_order(profit=-100)
        cancelled = self.make_order(status='cancelled', profit=99999)
        pending = self.make_order(status='cancel_refund_pending', profit=-99999)
        for token in ('profit', '-profit'):
            rows = list(sort_orders(SalesOrder.objects.all(), [token]))
            self.assertEqual(rows[0].pk, active.pk)
            self.assertEqual({row.pk for row in rows[1:]}, {cancelled.pk, pending.pk})
            self.assertTrue(all(row._sort_profit is None for row in rows[1:]))

    def test_sort_remove_and_reorder_preserve_other_conditions(self):
        context = sort_context(QueryDict('q=abc&status=completed&page=3&per_page=25'), ['-registration_date','owner_name','profit'])
        active = context['sort_active']
        def params(url):
            return parse_qs(urlsplit(url).query)
        self.assertEqual(params(active[1]['remove_url'])['sort'], ['-registration_date,profit'])
        self.assertEqual(params(active[1]['up_url'])['sort'], ['owner_name,-registration_date,profit'])
        self.assertEqual(params(active[1]['down_url'])['sort'], ['-registration_date,profit,owner_name'])
        for row in active:
            query = params(row['remove_url'])
            self.assertEqual(query['q'], ['abc'])
            self.assertEqual(query['per_page'], ['25'])
            self.assertNotIn('page', query)
        self.assertFalse(active[0]['up_url'])
        self.assertFalse(active[-1]['down_url'])
        self.assertEqual(active[0]['direction_label'], '新→舊')

    def test_dashboard_and_help_render_new_controls(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, '近 12 個月公司走勢')
        self.assertContains(response, '已有收款待確認')
        self.assertContains(response, '不是包含全公司費用的公司淨利')
        self.assertEqual(response.content.decode().count('viewBox="0 0 480 165"'), 3)
        response = self.client.get(reverse('order_list'), {'sort':'registration_date,owner_name'})
        self.assertContains(response, '移除車主排序')
        self.assertContains(response, '提前車主排序順位')
        self.assertContains(response, '全部清除')
        self.assertNotContains(response, '重新排序請先清除')
        self.assertContains(self.client.get(reverse('user_guide')), '新版戰情首頁、排序與狀態')

    def test_dashboard_grant_does_not_grant_financial_drilldown(self):
        from sales.access.models import ScreenAccessGrant, UserAccessState
        user = get_user_model().objects.create_user('dashboard-only', password='test-only')
        UserAccessState.objects.create(user=user, configured=True)
        ScreenAccessGrant.objects.create(user=user, screen_key='dashboard', view=True)
        self.client.force_login(user)
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, '近 12 個月公司走勢')
        self.assertNotContains(response, 'href="/operations/?')
        self.assertNotContains(response, '＋ 建立訂單')
        self.assertEqual(self.client.get(reverse('operations_report'), {'risk':'refund'}).status_code, 403)
        self.assertEqual(self.client.get(reverse('operations_report_export')).status_code, 403)
