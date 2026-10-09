"""舊車主資料（1.59.0）：Excel「舊車車主」的有／無是汰舊標記；舊車主電話與戶籍匯入、顯示；資料遷移可反向。"""
from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from sales.models import LegacyImportBatch, OrderOperationsProfile, SalesOrder, Store
from sales.services.legacy_import import build_import_preview, confirm_import, file_sha256
from sales.tests.test_legacy_import import workbook_bytes


class OldOwnerImportTests(TestCase):
    def setUp(self):
        Store.objects.create(name="總店", code="MAIN")
        self.tempdir = TemporaryDirectory()
        override = override_settings(MEDIA_ROOT=self.tempdir.name)
        override.enable()
        self.addCleanup(override.disable)
        self.addCleanup(self.tempdir.cleanup)

    def import_row(self, **cells):
        workbook = load_workbook(BytesIO(workbook_bytes()))
        for cell, value in cells.items():
            workbook["銷貨"][cell] = value
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("sales.xlsx", stream.getvalue())
        batch = LegacyImportBatch.objects.create(
            import_type="operations", source_file=upload, original_filename="sales.xlsx",
            file_sha256=file_sha256(upload), file_size=len(stream.getvalue()), uploaded_by="tester",
        )
        build_import_preview(batch, sheets=["銷貨"])
        confirm_import(batch, "tester")
        return SalesOrder.objects.get()

    def test_you_marker_means_trade_in_without_storing_a_name(self):
        order = self.import_row(BL4="有", BN4=None)
        self.assertTrue(order.is_trade_in_subsidy)
        self.assertEqual(order.old_owner_name, "")

    def test_wu_marker_means_no_trade_in(self):
        order = self.import_row(BL4="無", BM4="無", BN4=None)
        self.assertFalse(order.is_trade_in_subsidy)
        self.assertEqual((order.old_owner_name, order.old_owner_id_number), ("", ""))

    def test_old_owner_phone_and_household_are_imported(self):
        order = self.import_row(BL4="舊車主甲", BN4="OLD-123", BV4=912000111, BU4="臺北市中正區")
        self.assertTrue(order.is_trade_in_subsidy)
        self.assertEqual(order.old_owner_name, "舊車主甲")
        self.assertEqual(order.old_owner_phone, "912000111")
        self.assertEqual(order.old_owner_household_address, "臺北市中正區")


class OldOwnerDisplayTests(TestCase):
    def setUp(self):
        Store.objects.create(name="總店", code="MAIN")
        self.user = get_user_model().objects.create_superuser("admin", password="Test-Only-123")
        self.client.force_login(self.user)
        from sales.tests.profit_helpers import unlock_profit
        unlock_profit(self, self.user)
        tempdir = TemporaryDirectory()
        override = override_settings(MEDIA_ROOT=tempdir.name)
        override.enable()
        self.addCleanup(override.disable)
        self.addCleanup(tempdir.cleanup)
        workbook = load_workbook(BytesIO(workbook_bytes()))
        workbook["銷貨"]["BL4"], workbook["銷貨"]["BN4"] = "舊車主乙", "OLD-456"
        workbook["銷貨"]["BV4"], workbook["銷貨"]["BU4"] = "0922333444", "新北市板橋區"
        workbook["銷貨"]["BO4"], workbook["銷貨"]["BP4"] = "ENG-789", "三陽"
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("sales.xlsx", stream.getvalue())
        batch = LegacyImportBatch.objects.create(
            import_type="operations", source_file=upload, original_filename="sales.xlsx",
            file_sha256=file_sha256(upload), file_size=len(stream.getvalue()), uploaded_by="tester",
        )
        build_import_preview(batch, sheets=["銷貨"])
        confirm_import(batch, "tester")
        self.order = SalesOrder.objects.get()

    def test_subsidy_section_shows_contact_and_old_vehicle_data(self):
        html = self.client.get(reverse("order_detail", args=[self.order.pk])).content.decode()
        self.assertIn("舊車主電話", html)
        self.assertIn("0922333444", html)
        self.assertIn("新北市板橋區", html)
        self.assertIn("舊車與報廢資料", html)
        self.assertIn("ENG-789", html)
        self.assertEqual(OrderOperationsProfile.objects.get(order=self.order).old_vehicle_brand, "三陽")


class OldOwnerMigrationTests(TransactionTestCase):
    migrate_from = ("sales", "0168_sales_order_old_owner_contact")
    migrate_to = ("sales", "0169_old_owner_markers_and_contact_backfill")

    def setUp(self):
        super().setUp()
        MigrationExecutor(connection).migrate([self.migrate_from])
        self.addCleanup(self._restore)

    def _restore(self):
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())

    def _apps(self, target):
        return MigrationExecutor(connection).loader.project_state([target]).apps

    def test_markers_cleaned_backfilled_and_reversible(self):
        apps = self._apps(self.migrate_from)
        Store = apps.get_model("sales", "Store")
        Order = apps.get_model("sales", "SalesOrder")
        Batch = apps.get_model("sales", "LegacyImportBatch")
        Row = apps.get_model("sales", "LegacyImportRow")
        Snapshot = apps.get_model("sales", "LegacySalesSnapshot")
        Store.objects.get_or_create(code="MAIN", defaults={"name": "總店"})
        model = apps.get_model("sales", "VehicleModel").objects.create(
            brand="SUZUKI", name="舊車主遷移測試", model_number="OLD-OWNER-TEST", energy_type="gas", displacement_cc=125)
        color = apps.get_model("sales", "VehicleColor").objects.create(vehicle_model_id=model.pk, name="白")

        def order(name, **extra):
            return Order.objects.create(number=f"SO-OLD-OWNER-{name}", owner_name=name, owner_phone="0900", owner_address="地址",
                                        owner_id_number=f"HIST-{name}", vehicle_model_id=model.pk, color_id=color.pk,
                                        transaction_type="regular_new", vehicle_category=extra.pop("category", "new"), **extra)

        has = order("甲", old_owner_name="有", old_owner_id_number="有", is_trade_in_subsidy=False)
        none = order("乙", old_owner_name="無", is_trade_in_subsidy=False)
        flagged = order("丙", old_owner_name="有", is_trade_in_subsidy=True, trade_in_plate="ABC-1")
        used = order("丁", old_owner_name="有", category="used")
        real = order("戊", old_owner_name="真實舊車主", is_trade_in_subsidy=True)
        batch = Batch.objects.create(import_type="operations", source_file="x.xlsx", original_filename="x.xlsx",
                                     file_sha256="c" * 64, file_size=0, uploaded_by="t")
        row = Row.objects.create(batch=batch, sheet_name="銷貨", source_row=4, fingerprint="d" * 64, action="create",
                                 raw_data={"舊車車主電話": 912000111.0, "舊車戶籍": " 臺中市 "}, mapped_data={})
        Snapshot.objects.create(order=real, import_row=row)

        MigrationExecutor(connection).migrate([self.migrate_to])
        Order = self._apps(self.migrate_to).get_model("sales", "SalesOrder")
        OrderChange = self._apps(self.migrate_to).get_model("sales", "OrderChange")
        got = {o.owner_name: o for o in Order.objects.all()}
        self.assertEqual((got["甲"].old_owner_name, got["甲"].old_owner_id_number, got["甲"].is_trade_in_subsidy), ("", "", True))
        self.assertEqual((got["乙"].old_owner_name, got["乙"].is_trade_in_subsidy), ("", False))
        self.assertEqual((got["丙"].old_owner_name, got["丙"].is_trade_in_subsidy), ("", True))
        self.assertEqual((got["丁"].old_owner_name, got["丁"].is_trade_in_subsidy), ("", False), "中古車不標汰舊")
        self.assertEqual(got["戊"].old_owner_name, "真實舊車主")
        self.assertEqual((got["戊"].old_owner_phone, got["戊"].old_owner_household_address), ("912000111", "臺中市"))
        self.assertEqual(OrderChange.objects.count(), 4)

        MigrationExecutor(connection).migrate([self.migrate_from])
        Order = self._apps(self.migrate_from).get_model("sales", "SalesOrder")
        got = {o.owner_name: o for o in Order.objects.all()}
        self.assertEqual((got["甲"].old_owner_name, got["甲"].old_owner_id_number, got["甲"].is_trade_in_subsidy), ("有", "有", False))
        self.assertEqual(got["乙"].old_owner_name, "無")
        self.assertEqual((got["丙"].old_owner_name, got["丙"].is_trade_in_subsidy), ("有", True))
        self.assertEqual(self._apps(self.migrate_from).get_model("sales", "OrderChange").objects.count(), 0)
